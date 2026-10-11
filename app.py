from flask import Flask, render_template, request, session, redirect, send_file,send_from_directory
import os
import time
from docx import Document
from pptx import Presentation
import json
import sqlite3
from datetime import datetime, timezone, timedelta
from werkzeug.security import generate_password_hash, check_password_hash

from pypdf import PdfReader
from dotenv import load_dotenv
from google import genai
from google.genai import types

from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# Load API key
load_dotenv()

api_key = os.getenv("GEMINI_API_KEY")

if not api_key:
    raise ValueError("GEMINI_API_KEY was not found in the .env file.")

client = genai.Client(api_key=api_key)


# Flask application
app = Flask(__name__)

# Secret key for storing quiz data
app.secret_key = "ai_notes_quiz_secret_key"
def get_database():
    connection = sqlite3.connect("database.db")
    connection.row_factory = sqlite3.Row
    return connection


UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# -------------------------
# Home page
# -------------------------

@app.route("/")
def home():
    return render_template("index.html")
@app.route("/signup", methods=["GET", "POST"])
def signup():

    if request.method == "POST":

        username = request.form["username"]
        email = request.form["email"]
        password = request.form["password"]
        password = generate_password_hash(password)



        connection = get_database()
        cursor = connection.cursor()

        try:
            cursor.execute(
                """
                INSERT INTO users (username, email, password)
                VALUES (?, ?, ?)
                """,
                (username, email, password)
            )

            connection.commit()

        except sqlite3.IntegrityError:
            connection.close()
            return "Email already registered."

        connection.close()

        return "Account created successfully! <a href='/login'>Login here</a>"

    return render_template("signup.html")
@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        connection = get_database()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT * FROM users
            WHERE email = ?
            """,
            (email,)
        )

        user = cursor.fetchone()

        connection.close()

        if user and check_password_hash(user["password"], password):

            session["user_id"] = user["id"]
            session["username"] = user["username"]

            return redirect("/dashboard")

        else:

            return """
            <h1>Login Failed</h1>
            <p>Invalid email or password.</p>
            <a href="/login">Try Again</a>
            """

    return render_template("login.html")  
@app.route("/dashboard")
def dashboard():
    if "user_id" not in session:
        return """
        <h1>Please login first</h1>
        <a href="/login">Login</a>
        """

    connection = get_database()

    cursor = connection.cursor()

    cursor.execute(
        """
        SELECT score, total, created_at
        FROM quiz_history
        WHERE user_id = ?
        ORDER BY created_at DESC
        """,
        (session["user_id"],)
    )

    history = cursor.fetchall()
    total_quizzes = len(history)

    if history:
        total_score = sum(quiz["score"] for quiz in history)
        total_questions = sum(quiz["total"] for quiz in history)

        average_score = round(
            (total_score / total_questions) * 100,
            2
        )

        highest_score = max(
            quiz["score"] for quiz in history
        )
    else:
        total_score = 0
        total_questions = 0
        average_score = 0
        highest_score = 0
    connection.close()

    formatted_history = []

    for quiz in history:
            utc_time = datetime.strptime(
                quiz["created_at"],
                "%Y-%m-%d %H:%M:%S"
            ).replace(tzinfo=timezone.utc)

            ist_time = utc_time.astimezone(
                timezone(timedelta(hours=5, minutes=30))
            )

            formatted_date = ist_time.strftime(
             "%d %B %Y, %I:%M %p"
            )

            formatted_history.append({
                "score": quiz["score"],
                "total": quiz["total"],
                "created_at": formatted_date
            })

    return render_template(
    "dashboard.html",
    username=session["username"],
    history=formatted_history,
    total_quizzes=total_quizzes,
    total_questions=total_questions,
    average_score=average_score,
    highest_score=highest_score
    )
@app.route("/history")
def history():

    if "user_id" not in session:
        return redirect("/login")

    connection = get_database()
    cursor = connection.cursor()

    cursor.execute(
        """
        SELECT score, total, created_at
        FROM quiz_history
        WHERE user_id = ?
        ORDER BY created_at DESC
        """,
        (session["user_id"],)
    )

    history = cursor.fetchall()
    total_attempts = len(history)

    if history:
        total_questions = sum(quiz["total"] for quiz in history)
        total_score = sum(quiz["score"] for quiz in history)

        average_percentage = round(
            (total_score / total_questions) * 100,
            2
         )

        best_score = max(
            quiz["score"] for quiz in history
        )
    else:
        total_questions = 0
        total_score = 0
        average_percentage = 0
        best_score = 0
    connection.close()

    formatted_history = []

    for quiz in history:

        utc_time = datetime.strptime(
            quiz["created_at"],
            "%Y-%m-%d %H:%M:%S"
        ).replace(tzinfo=timezone.utc)

        ist_time = utc_time.astimezone(
            timezone(timedelta(hours=5, minutes=30))
        )

        formatted_date = ist_time.strftime(
            "%d %B %Y, %I:%M %p"
        )

        formatted_history.append({
            "score": quiz["score"],
            "total": quiz["total"],
            "created_at": formatted_date
        })

    return render_template(
    "history.html",
    history=formatted_history,
    total_attempts=total_attempts,
    average_percentage=average_percentage,
    best_score=best_score
)
@app.route("/logout")
def logout():

    session.clear()

    return """
    <h1>Logged out successfully</h1>
    <a href="/login">Login again</a>
    """

# -------------------------
# PDF upload
# -------------------------
@app.route("/upload", methods=["POST"])
def upload():
    if "user_id" not in session:
        return redirect("/login")
    difficulty = request.form.get("difficulty", "Medium")
    question_count = request.form.get("question_count", "5")
    print("Selected question count:", question_count)
    file = request.files["pdf_file"]

    if not file or file.filename == "":
        return "Please select a file."

    filename = file.filename.lower()

    if not (
        filename.endswith(".pdf")
        or filename.endswith(".docx")
        or filename.endswith(".txt")
        or filename.endswith(".pptx")
        or filename.endswith(".jpg")
        or filename.endswith(".jpeg")
        or filename.endswith(".png")
    ):
        return "Please upload a PDF, DOCX, TXT, PPTX, JPG, JPEG, or PNG file."

    # Save file
    file_path = os.path.join(
        app.config["UPLOAD_FOLDER"],
        file.filename
    )

    file.save(file_path)
    # Store uploaded image name
    if filename.endswith((".jpg", ".jpeg", ".png")):
         session["image_filename"] = file.filename
    else:
        session["image_filename"] = None

    # -------------------------
    # Read uploaded file
    # -------------------------

    text = ""

    if filename.endswith(".pdf"):

        reader = PdfReader(file_path)

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"


    elif filename.endswith(".docx"):

        document = Document(file_path)

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text += paragraph.text + "\n"


    elif filename.endswith(".txt"):

        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()


    elif filename.endswith(".pptx"):

        presentation = Presentation(file_path)

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text") and shape.text.strip():
                    text += shape.text + "\n"


    elif (
        filename.endswith(".jpg")
        or filename.endswith(".jpeg")
        or filename.endswith(".png")
    ):

        image_file = client.files.upload(file=file_path)

        image_prompt = """
        Analyze this image carefully.

        Extract all useful educational information from the image.
        If it contains a chart, explain the chart, values, labels,
        comparisons, and trends.
        If it contains a diagram, explain its important parts
        and relationships.
        If it contains text, extract the important text.

        Return the information as plain text suitable for
        generating college study notes and quiz questions.
        """

        image_response = client.models.generate_content(
            model="gemini-3.6-flash",
            contents=[image_file, image_prompt]
        )

        text = image_response.text


    if not text.strip():
        return "Could not extract text from this file."

    
    # -------------------------
    # Generate Notes
    # -------------------------

    prompt = f"""
Create concise study notes from the complete text.

Include only the most important points, definitions, key concepts,
advantages, disadvantages, important examples, and exam-relevant information.

Do not include unnecessary explanations or repeat the same information.
{text}

Use simple language suitable for a college student.

Format the notes with clear headings and bullet points.
Keep each point short and easy to revise.uage suitable for a college student.
"""

    notes = None

    for attempt in range(3):

        try:

            response = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=prompt
            )

            notes = response.text
            break

        except Exception as e:

            print("Gemini error:", e)

            if attempt < 2:

                wait_time = 5 * (2 ** attempt)

                print(
                    f"Retrying in {wait_time} seconds..."
                )

                time.sleep(wait_time)

            else:

                print("Gemini quota exhausted. Creating local study notes.")

                import re

                # Split the complete uploaded material into sentences
                sentences = re.split(r'(?<=[.!?])\s+', text.strip())

                # Remove very short sentences
                sentences = [
                    s.strip()
                    for s in sentences
                    if len(s.strip()) > 30
                ]

                # Keep important-looking sentences
                important_sentences = []

                keywords = [
                    "important",
                    "definition",
                    "define",
                    "advantage",
                    "advantages",
                    "disadvantage",
                    "disadvantages",
                    "purpose",
                    "types",
                    "features",
                    "function",
                    "process",
                    "example",
                    "uses",
                    "application",
                    "benefit"
                ]

                for sentence in sentences:

                    lower_sentence = sentence.lower()

                    if any(keyword in lower_sentence for keyword in keywords):
                        important_sentences.append(sentence)

                # If not enough important sentences, use the first sentences
                if len(important_sentences) < 5:
                    important_sentences = sentences

                # Remove duplicate sentences
                important_sentences = list(dict.fromkeys(important_sentences))

                # Limit the number of points
                important_sentences = important_sentences[:20]

                notes = "# Study Notes\n\n"

                notes += "## Key Points\n\n"

                for sentence in important_sentences:
                    notes += "- " + sentence + "\n"

                notes += "\n## Summary\n\n"

                summary_sentences = sentences[:5]

                for sentence in summary_sentences:
                    notes += sentence + " "

                notes += "\n\n## Revision Tip\n\n"
                notes += "Review the key points first and then read the summary for quick revision."

                print("Local study notes created successfully.")
    # -------------------------
    # Generate AI Quiz
    # -------------------------

    time.sleep(5)

    quiz_prompt = f"""
    Create exactly {question_count} multiple-choice questions from the study material below.

    Quiz difficulty level: {difficulty}

    Return ONLY valid JSON.

    The JSON must have this format:

    {{
        "questions": [
                {{
                    "question": "Question text",
                    "options": {{
                        "A": "Option A",
                        "B": "Option B",
                        "C": "Option C",
                        "D": "Option D"
                    }},
                    "answer": "A"
                }}
            ]
    }}

        Rules:
       - Create exactly {question_count} questions.
        - Each question must have exactly 4 options.
        - Options must be A, B, C and D.
        - The answer must be only A, B, C or D.
        - The answer must be the correct option.
        - Questions must be based only on the study material.
        - Use simple language suitable for a college student.
        - Do not include Markdown.
        - Do not include explanations.
        - Do not include any text outside the JSON.

        Difficulty rules:
        - Easy: Ask basic questions that test simple understanding and important facts.
        - Medium: Ask questions that require understanding and applying concepts.
        - Hard: Ask questions that require deeper understanding, comparison, reasoning, or application.

        Study material:

        {text[:1000]}
        """

    quiz_data = None
    quiz_response = None

    for attempt in range(3):

        try:

            quiz_response = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=quiz_prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                )
            )

            quiz_data = json.loads(quiz_response.text)

            session["quiz"] = quiz_data["questions"]

            print("AI quiz generated successfully.")

            break

        except Exception as e:

            print("Quiz generation error:", e)
            if "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e):

                print("Gemini quota exhausted. Using uploaded material.")

                import re

                sentences = re.split(
                    r'(?<=[.!?])\s+',
                    text.strip()
                )

                sentences = [
                    s.strip()
                    for s in sentences
                    if len(s.strip()) > 40
                ]

                quiz_questions = []

                related_words = [
                    "synchronization",
                    "communication",
                    "process",
                    "thread",
                    "memory",
                    "resource",
                    "operation",
                    "management",
                    "system",
                    "execution"
                ]

                for sentence in sentences[:int(question_count)]:

                    words = sentence.split()

                    if len(words) < 8:
                        continue

                    index = min(4, len(words) - 1)

                    answer_word = re.sub(
                        r'[^\w]',
                        '',
                        words[index]
                    )

                    if not answer_word:
                        continue

                    question_text = sentence.replace(
                        words[index],
                        "_____",
                        1
                    )

                    wrong_options = [
                        word for word in related_words
                        if word.lower() != answer_word.lower()
                    ]

                    options = {
                        "A": answer_word,
                        "B": wrong_options[0],
                        "C": wrong_options[1],
                        "D": wrong_options[2]
                    }

                    quiz_questions.append({
                        "question": "Fill in the blank: " + question_text,
                        "options": options,
                        "answer": "A"
                    })

                if not quiz_questions:

                    quiz_questions = [
                        {
                            "question": "Which topic is discussed in the uploaded study material?",
                            "options": {
                                "A": "Synchronization",
                                "B": "Database management",
                                "C": "Computer graphics",
                                "D": "Web development"
                            },
                            "answer": "A"
                        }
                    ]

                quiz_data = {
                    "questions": quiz_questions
                }

                session["quiz"] = quiz_data["questions"]

                print("Local quiz created from uploaded material.")

                break
            
             
            elif attempt < 2:

                wait_time = 15 * (2 ** attempt)

                print(
                    f"Retrying quiz generation in {wait_time} seconds..."
                )

                time.sleep(wait_time)

            else:

                quiz_data = {
                    "questions": []
                }

                session["quiz"] = []

                quiz_error = "⚠️ Quiz generation is temporarily unavailable."


    # -------------------------
    # Display Notes and Quiz
    # -------------------------

    return render_template(
        "text.html",
        text=text,
        notes=notes,
        quiz=quiz_data["questions"] if quiz_data and quiz_data["questions"] else quiz_error
    )
@app.route("/uploads/<filename>")
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)    
# -------------------------
# Interactive AI Quiz
# -------------------------

@app.route("/quiz")
def quiz():

    questions = session.get("quiz", [])

    if not questions:
        return """
        <h1>No quiz available</h1>
        <p>Please upload a PDF first to generate an AI quiz.</p>
        <a href="/">Upload PDF</a>
        """

    return render_template(
        "quiz.html",
        questions=questions
    )


# -------------------------
# Submit AI Quiz
# -------------------------

@app.route("/submit_quiz", methods=["POST"])
def submit_quiz():
    if "user_id" not in session:
        return redirect("/login")

    questions = session.get("quiz", [])

    if not questions:
        return "Quiz data not found. Please generate a new quiz."

    score = 0

    results = []

    for index, question in enumerate(questions):

        question_number = f"q{index + 1}"

        user_answer = request.form.get(question_number)

        correct_answer = question["answer"]

        if user_answer == correct_answer:
            score += 1
            status = "Correct"
        else:
            status = "Wrong"

        results.append({
            "question": question["question"],
            "user_answer": user_answer,
            "correct_answer": correct_answer,
            "status": status
        })


    connection = get_database()

    cursor = connection.cursor()

    cursor.execute(
        """
        INSERT INTO quiz_history (user_id, score, total)
        VALUES (?, ?, ?)
        """,
        (
            session["user_id"],
            score,
            len(questions)
        )
    )

    connection.commit()
    connection.close()

    return render_template(
        "result.html",
        score=score,
        total=len(questions),
        results=results
    )
@app.route("/delete_history")
def delete_history():

    if "user_id" not in session:
        return redirect("/login")

    connection = get_database()
    cursor = connection.cursor()

    cursor.execute(
        """
        DELETE FROM quiz_history
        WHERE user_id = ?
        """,
        (session["user_id"],)
    )

    connection.commit()
    connection.close()

    return redirect("/history")

@app.route("/download_notes")
def download_notes():

    if "user_id" not in session:
        return redirect("/login")

    notes = session.get("notes", "")

    if not notes:
        return "Notes not available. Please generate notes again."

    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=A4)

    width, height = A4

    x = 40
    y = height - 50

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(x, y, "AI Generated Study Notes")

    y -= 30

    pdf.setFont("Helvetica", 11)

    for line in notes.splitlines():

        if y < 50:
            pdf.showPage()
            pdf.setFont("Helvetica", 11)
            y = height - 50

        pdf.drawString(x, y, line[:100])

        y -= 18

    pdf.save()

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="AI_Study_Notes.pdf",
        mimetype="application/pdf"
    )    
# -------------------------
# Start Flask
# -------------------------

if __name__ == "__main__":
    app.run(debug=True)